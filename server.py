import base64
import io
import json
import math
import os
import re
import sqlite3
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from groq import Groq
from openpyxl import load_workbook

ROOT = Path(__file__).parent
DATABASE = sqlite3.connect(":memory:", check_same_thread=False)
DATABASE.row_factory = sqlite3.Row
MODEL = "openai/gpt-oss-20b"
WORKBOOKS = []
ACTIVITY = []


def load_local_environment():
    env_file = ROOT / ".env"
    if not env_file.exists():
        return
    for line in env_file.read_text(encoding="utf-8").splitlines():
        key, separator, value = line.partition("=")
        if separator and key.strip():
            os.environ[key.strip()] = value.strip().strip('"').strip("'")


load_local_environment()


def safe_table_name(name):
    cleaned = re.sub(r"[^a-zA-Z0-9_]", "_", name.strip().lower()).strip("_")
    return cleaned or "sheet"


def load_workbook_bytes(file_bytes):
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    imported = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = []
        for index, value in enumerate(rows[0]):
            header = safe_table_name(str(value or f"column_{index + 1}"))
            if header in headers:
                header = f"{header}_{index + 1}"
            headers.append(header)
        table = safe_table_name(sheet.title)
        DATABASE.execute(f'DROP TABLE IF EXISTS "{table}"')
        columns = ", ".join(f'"{header}" TEXT' for header in headers)
        DATABASE.execute(f'CREATE TABLE "{table}" ({columns})')
        placeholders = ",".join("?" for _ in headers)
        for row in rows[1:]:
            values = [None if value is None else str(value) for value in row[: len(headers)]]
            values += [None] * (len(headers) - len(values))
            DATABASE.execute(f'INSERT INTO "{table}" VALUES ({placeholders})', values)
        imported.append({"sheet": sheet.title, "table": table, "rows": max(len(rows) - 1, 0)})
    DATABASE.commit()
    return imported


def _quote_identifier(name):
    # Only allow identifiers that already match our safe_table_name format.
    if not re.match(r"^[a-zA-Z0-9_]+$", name or ""):
        raise ValueError(f"Invalid identifier: {name!r}")
    return f'"{name}"'


def execute_read_only_sql(query):
    normalized = query.strip()
    if not re.match(r"^(select|with)\b", normalized, re.IGNORECASE):
        raise ValueError("Only SELECT or WITH queries are allowed")
    if ";" in normalized.rstrip(";"):
        raise ValueError("Multiple SQL statements are not allowed")
    if re.search(r"\b(insert|update|delete|drop|alter|create|attach|pragma|vacuum)\b", normalized, re.IGNORECASE):
        raise ValueError("Only read-only SQL is allowed")
    cursor = DATABASE.execute(normalized)
    columns = [description[0] for description in cursor.description or []]
    rows = [dict(zip(columns, row)) for row in cursor.fetchmany(100)]
    return {"query": query, "columns": columns, "rows": rows, "row_count": len(rows)}


def _to_float(value):
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def project_metric(table, value_column, periods_ahead=3, date_column=None, method="linear"):
    """
    Pulls a historical numeric series out of an imported sheet and projects it
    forward. Does the actual math in Python (not via the LLM) so projections
    are reproducible and auditable, rather than the model guessing numbers.
    """
    table_id = _quote_identifier(table)
    value_id = _quote_identifier(value_column)
    periods_ahead = max(1, min(int(periods_ahead), 24))

    if date_column:
        date_id = _quote_identifier(date_column)
        query = f"SELECT {date_id} AS x_label, {value_id} AS y_value FROM {table_id} ORDER BY {date_id}"
    else:
        query = f"SELECT {value_id} AS y_value FROM {table_id}"

    cursor = DATABASE.execute(query)
    raw_rows = cursor.fetchall()

    series = []
    for row in raw_rows:
        y = _to_float(row["y_value"])
        if y is None:
            continue
        label = row["x_label"] if date_column and "x_label" in row.keys() else None
        series.append((label, y))

    if len(series) < 2:
        return {
            "error": "Not enough historical numeric data points to project a trend "
                     f"(found {len(series)}). Need at least 2.",
            "table": table,
            "value_column": value_column,
        }

    ys = [point[1] for point in series]
    n = len(ys)
    xs = list(range(n))

    mean_x = sum(xs) / n
    mean_y = sum(ys) / n
    numerator = sum((xs[i] - mean_x) * (ys[i] - mean_y) for i in range(n))
    denominator = sum((xs[i] - mean_x) ** 2 for i in range(n))
    slope = numerator / denominator if denominator else 0.0
    intercept = mean_y - slope * mean_x

    linear_forecast = []
    for step in range(periods_ahead):
        x = n + step
        linear_forecast.append(round(intercept + slope * x, 2))

    cagr = None
    cagr_forecast = []
    first, last = ys[0], ys[-1]
    if first > 0 and last > 0 and n > 1:
        periods_elapsed = n - 1
        cagr = (last / first) ** (1 / periods_elapsed) - 1
        for step in range(1, periods_ahead + 1):
            cagr_forecast.append(round(last * ((1 + cagr) ** step), 2))

    chosen_method = method if method in ("linear", "cagr") else "linear"
    forecast_values = linear_forecast if chosen_method == "linear" else (cagr_forecast or linear_forecast)

    return {
        "table": table,
        "value_column": value_column,
        "date_column": date_column,
        "historical_values": ys,
        "historical_labels": [point[0] for point in series] if date_column else None,
        "method_used": chosen_method,
        "linear_trend": {"slope_per_period": round(slope, 4), "intercept": round(intercept, 4)},
        "compound_growth_rate_per_period": round(cagr, 4) if cagr is not None else None,
        "projected_values": forecast_values,
        "periods_ahead": periods_ahead,
        "note": "Estimate derived purely from historical trend in the imported data. "
                "Not a guarantee of future performance.",
    }


def describe_schema():
    """
    Builds a human-readable summary of every imported table, its columns, and
    - for low-cardinality text columns (categories, vendors, statuses, modes,
    etc.) - the actual distinct values present. This lets the model map fuzzy
    business concepts ("fixed expenses", "recurring costs") onto the literal
    strings that exist in the data (e.g. "Rent", "Payroll", "AWS") instead of
    searching for a column or value that was never going to exist.

    High-cardinality columns (free text, IDs, raw numbers) are intentionally
    left without value samples - listing every distinct amount or reference
    number would bloat the prompt without helping the model reason.
    """
    tables = DATABASE.execute(
        "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
    ).fetchall()
    if not tables:
        return "No workbook has been imported yet - no tables are available."

    DISTINCT_VALUE_LIMIT = 30  # only sample columns with at most this many distinct values
    VALUE_PREVIEW_MAX_CHARS = 400

    lines = []
    for row in tables:
        table = row["name"]
        columns = DATABASE.execute(f'PRAGMA table_info("{table}")').fetchall()
        row_count = DATABASE.execute(f'SELECT COUNT(*) AS c FROM "{table}"').fetchone()["c"]
        lines.append(f"- {table} ({row_count} rows):")

        for col in columns:
            col_name = col["name"]
            col_id = _quote_identifier(col_name)
            distinct_rows = DATABASE.execute(
                f'SELECT DISTINCT {col_id} AS v FROM "{table}" '
                f'WHERE {col_id} IS NOT NULL LIMIT {DISTINCT_VALUE_LIMIT + 1}'
            ).fetchall()
            distinct_values = [r["v"] for r in distinct_rows]

            # Only show a value sample when the column looks categorical: a
            # small, bounded set of repeating text values (not free text,
            # not a raw number/amount, not a unique ID/reference column).
            looks_categorical = (
                0 < len(distinct_values) <= DISTINCT_VALUE_LIMIT
                and len(distinct_values) < row_count
                and not all(_to_float(v) is not None for v in distinct_values)
            )

            if looks_categorical:
                preview = ", ".join(str(v) for v in distinct_values)
                if len(preview) > VALUE_PREVIEW_MAX_CHARS:
                    preview = preview[:VALUE_PREVIEW_MAX_CHARS] + ", ..."
                lines.append(f'    "{col_name}" - values seen: {preview}')
            else:
                lines.append(f'    "{col_name}"')

    return "\n".join(lines)


def describe_available_data():
    tables = DATABASE.execute(
        "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name"
    ).fetchall()
    if not tables:
        return "No workbook data is imported yet. Upload an Excel workbook to get started."

    descriptions = []
    for row in tables:
        table = row["name"]
        columns = DATABASE.execute(f'PRAGMA table_info("{table}")').fetchall()
        row_count = DATABASE.execute(
            f'SELECT COUNT(*) AS c FROM "{table}"'
        ).fetchone()["c"]
        column_names = ", ".join(column["name"] for column in columns)
        descriptions.append(f"- **{table}**: {row_count} rows ({column_names})")
    return "I have the following data from your imported workbook:\n\n" + "\n".join(descriptions)


def call_groq(question, session_summary):
    if re.search(
        r"\b(what data|what do i have|what(?:'s| is) in (?:the )?(?:workbook|excel|data)|available data|what sheets?)\b",
        question,
        re.IGNORECASE,
    ):
        return {"answer": describe_available_data(), "evidence": None}

    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        raise RuntimeError("GROQ_API_KEY is not configured. Add it to the environment before asking a question.")

    tools = [
        {
            "type": "function",
            "function": {
                "name": "execute_sql",
                "description": "Run a read-only SELECT or WITH query against imported Excel sheets. "
                                "Use this whenever the user asks for a specific figure, comparison, "
                                "or fact that should come from the workbook.",
                "parameters": {
                    "type": "object",
                    "properties": {"query": {"type": "string"}},
                    "required": ["query"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "custom_sql_query",
                "description": "Custom fallback: write a read-only SELECT or WITH query against any imported "
                                "Excel sheet when no narrower tool matches.",
                "parameters": {
                    "type": "object",
                    "properties": {"query": {"type": "string"}, "reason": {"type": "string"}},
                    "required": ["query", "reason"],
                },
            },
        },
        {
            "type": "function",
            "function": {
                "name": "project_metric",
                "description": "Forecast/project a numeric column forward using its historical values in an "
                                "imported sheet. Use this for ANY question asking about future, expected, "
                                "forecasted, or projected numbers (e.g. 'what will revenue be next quarter'). "
                                "Do not attempt to calculate projections yourself - always call this tool so "
                                "the math is verifiable.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "table": {"type": "string", "description": "Table name to pull historical data from."},
                        "value_column": {"type": "string", "description": "Numeric column to project."},
                        "date_column": {
                            "type": "string",
                            "description": "Optional column to order the series by (e.g. a date/period column).",
                        },
                        "periods_ahead": {
                            "type": "integer",
                            "description": "How many future periods to project. Default 3.",
                        },
                        "method": {
                            "type": "string",
                            "enum": ["linear", "cagr"],
                            "description": "Projection method: 'linear' trend or 'cagr' compound growth. "
                                            "Default linear.",
                        },
                    },
                    "required": ["table", "value_column"],
                },
            },
        },
    ]

    system = f"""
You are Ledgerly, an intelligent business assistant that answers questions about the user's business workbook.

You should behave like a practical financial/business analyst: understand natural language, make reasonable interpretations, perform calculations when appropriate, and use the workbook whenever the question depends on the user's data.

========================
1. GENERAL BUSINESS QUESTIONS
========================

For general business, finance, accounting, strategy, or educational questions that do NOT require the user's workbook data, answer directly from your own knowledge.

Examples:
- What is gross profit?
- What is working capital?
- How can I reduce operating expenses?
- What is the difference between revenue and profit?

Do not call a tool for these questions.

========================
2. WORKBOOK DATA QUESTIONS
========================

Whenever the user asks for a specific fact, number, amount, comparison, trend, transaction, category, or historical value from their workbook, you MUST call execute_sql or custom_sql_query first.

Examples:
- How much did I spend in January?
- What was my revenue in Q1?
- Which vendor did I pay the most?
- Compare Q1 and Q2 expenses.
- How much did I spend on AWS?

Rules:
- Never invent a workbook value.
- Only use numbers returned by SQL.
- If SQL returns no rows, say that the workbook has no evidence for the requested information.
- Always mention the sheet/table used for the answer.

========================
3. DERIVED METRICS
========================

Some business metrics are not stored directly and must be calculated.

Examples:
- Profit = Revenue - Expenses
- Net Cash Flow = Money In - Money Out
- Growth % = ((Current Period - Previous Period) / Previous Period) * 100
- Average Monthly Expense = Total Expenses / Number of Months
- Average Transaction Value = Total Amount / Number of Transactions

For derived metrics:
1. Query the raw workbook values using SQL.
2. Calculate the metric from the returned values.
3. Clearly present it as a calculated/derived result.

Only use table names and column names that literally exist in the schema below.

========================
4. FUTURE VALUES / FORECASTS
========================

For ANY question about a future, expected, projected, predicted, likely, estimated, or approximate FUTURE value, use project_metric.

Examples:
- What will my expenses be next month?
- How much revenue will I make next quarter?
- What should I expect my expenses to be?
- How much am I likely to spend?
- Roughly how much will I spend next month?
- Forecast my revenue.

NEVER manually invent a future number.

The projected number must come from project_metric.

Always describe the returned value as an estimate or projection, NOT a guaranteed result.

Always mention the method used:
- Linear trend
- Compound growth rate

Example:
"Your estimated expenses for next month are approximately ₹85,000, based on the historical trend using a linear trend model."

If there is not enough historical data, clearly say that a reliable projection cannot be produced.

========================
5. IMPORTANT: HISTORICAL APPROXIMATION VS FUTURE ESTIMATION
========================

The word "roughly", "approximately", or "guess" does NOT automatically mean future forecasting.

For HISTORICAL questions:
- Use SQL.
- Calculate or summarize from actual workbook data.

Examples:
- Roughly how much did I spend last year?
- Approximately what were my Q1 expenses?
- Give me an approximate average monthly spend.

For FUTURE questions:
- Use project_metric.

Examples:
- Roughly how much will I spend next month?
- Approximately what will my revenue be next quarter?

========================
6. TIME PERIOD UNDERSTANDING
========================

Understand common business time expressions automatically.

Interpret:

Today = current date
Yesterday = previous calendar day
This week = current calendar week
Last week = previous calendar week
This month = current calendar month
Last month = previous calendar month
This quarter = current quarter
Last quarter = previous quarter
This year = current calendar year
Last year = previous calendar year

========================
7. QUARTERS
========================

Understand Q1, Q2, Q3, and Q4 automatically.

Default calendar-year interpretation:

Q1 = January through March
Q2 = April through June
Q3 = July through September
Q4 = October through December

If the workbook clearly uses a fiscal year, use the fiscal-year convention instead.

Example for an April-March fiscal year:

Q1 = April-June
Q2 = July-September
Q3 = October-December
Q4 = January-March

Do not ask the user what Q1 means unless the available workbook/context makes it genuinely ambiguous.

Examples:

"Q1 expenses"
→ Query expenses for the applicable Q1 period.

"Q2 revenue"
→ Query revenue for the applicable Q2 period.

"Compare Q1 and Q2"
→ Query both periods and compare them.

If no year is specified:
- Prefer the year implied by context or workbook data.
- If multiple years exist and the question is ambiguous, use the most recent applicable period and briefly state the assumption.

========================
8. NATURAL DATE LANGUAGE
========================

Understand phrases such as:

"previous 3 months"
"last 6 months"
"past year"
"this quarter"
"last quarter"
"next month"
"next quarter"
"year to date"
"YTD"
"Q1"
"Q2"
"Q3"
"Q4"

Convert them into appropriate date ranges or periods before generating SQL.

========================
9. FUZZY BUSINESS CONCEPTS
========================

Users may ask about concepts that do not literally exist as categories in the workbook.

Examples:
- Fixed expenses
- Recurring costs
- Variable expenses
- Discretionary spending
- Essential spending
- Overhead
- Operating expenses
- Marketing costs

For these questions:

1. Inspect the "values seen" information in the schema.
2. Identify which EXISTING literal values reasonably match the user's concept.
3. Generate SQL using only those existing literal values.
4. Briefly tell the user which values were grouped under the concept.

Example:
"I treated Rent, Payroll, and Subscriptions as fixed/recurring expenses based on the categories present in the workbook."

Important:
- Never invent a category.
- Never claim a value exists unless it appears in the schema.
- The classification itself may be an inference, but the underlying values must exist.

========================
10. NATURAL BUSINESS LANGUAGE
========================

Interpret ordinary business language intelligently.

Examples:

"What am I spending?"
→ Interpret as expenses/outgoing money.

"How much am I making?"
→ Interpret as revenue/income.

"How much did I burn?"
→ Interpret based on context as burn rate or net cash outflow.

"What are my recurring costs?"
→ Identify recurring/fixed-like categories from workbook values.

"How much did I spend?"
→ Historical workbook question → SQL.

"How much will I spend?"
→ Future question → project_metric.

"What should I expect?"
→ Future estimate → project_metric.

Do not unnecessarily ask the user to rephrase normal business language.

========================
11. COMPARISONS
========================

For comparisons involving workbook data, use SQL first.

Examples:
- Was Q2 more expensive than Q1?
- Which month had the highest revenue?
- Did expenses increase?
- Compare this year's revenue with last year.

Return:
- Relevant values
- Difference
- Percentage change when useful
- Clear conclusion

========================
12. FACT VS CALCULATION VS INFERENCE VS FORECAST
========================

Keep these categories separate.

FACT:
"February expenses were ₹72,400."
→ Must come directly from SQL.

CALCULATED:
"February profit was ₹31,600."
→ Calculated from SQL-returned values.

INFERENCE:
"I classified Rent and Payroll as fixed expenses."
→ Business classification based on actual workbook categories.

FORECAST:
"Next month's expenses are estimated at ₹75,000."
→ Must come from project_metric.

Never present:
- An inference as a fact.
- A calculation as a stored workbook value.
- A forecast as a historical value.

========================
13. ASSUMPTIONS
========================

You may make reasonable assumptions when interpreting:
- Q1/Q2/Q3/Q4
- Calendar periods
- Fiscal periods
- Common business terminology
- Natural-language dates
- Approximate requests

Only mention an assumption when it materially affects the answer.

Example:
"Assuming Q1 refers to January-March..."

========================
14. TOOL SELECTION
========================

Use NO TOOL for:
- General knowledge
- Definitions
- General business advice
- General strategy
- General explanations

Use execute_sql or custom_sql_query for:
- Historical workbook data
- Specific workbook values
- Comparisons
- Aggregations
- Derived metrics
- Historical approximations
- Fuzzy business concepts

Use project_metric for:
- Future values
- Forecasts
- Predictions
- Expected values
- Next month
- Next quarter
- Next year
- Future estimates

========================
15. DATA ACCURACY
========================

Never fabricate workbook facts.

Never guess a historical number.

Never invent a missing category.

Never use a table or column that does not literally exist in the schema.

Never present a forecast as an actual value.

If the workbook does not contain enough evidence, say so.

========================
16. AVAILABLE TABLES AND COLUMNS
========================

This is the complete schema. Nothing else exists:

{describe_schema()}

Only use table names and column names that literally exist in this schema.

========================
17. CURRENT SESSION CONTEXT
========================

{session_summary or "No prior context in this session."}

Use this only as conversational context. Workbook facts must still be verified using the appropriate tool.

========================
18. FINAL RESPONSE STYLE
========================

Be concise, natural, and business-focused.

Do not expose internal reasoning or chain-of-thought.

For workbook answers:
- Give the answer first.
- Mention the source sheet/table.
- Include calculations when useful.

For forecasts:
- Clearly say "estimated", "projected", or "approximately".
- Mention the forecasting method.

For fuzzy classifications:
- Briefly state the categories used.

Ledgerly should feel like an intelligent financial analyst: flexible in understanding what the user means, but strict about not fabricating actual workbook data.
"""

    client = Groq(api_key=api_key)
    messages = [{"role": "system", "content": system}, {"role": "user", "content": question}]
    last_evidence = None

    tool_impls = {
        "execute_sql": lambda args: execute_read_only_sql(args["query"]),
        "custom_sql_query": lambda args: execute_read_only_sql(args["query"]),
        "project_metric": lambda args: project_metric(
            table=args["table"],
            value_column=args["value_column"],
            periods_ahead=args.get("periods_ahead", 3),
            date_column=args.get("date_column"),
            method=args.get("method", "linear"),
        ),
    }

    for _ in range(5):
        completion = client.chat.completions.create(
            model=MODEL,
            messages=messages,
            tools=tools,
            tool_choice="auto",
            temperature=1,
            max_completion_tokens=2048,
            top_p=1,
            reasoning_effort="medium",
            stream=False,
        )
        message = completion.choices[0].message
        tool_calls = message.tool_calls or []

        if not tool_calls:
            # No tool needed (e.g. general business knowledge question) - answer directly.
            return {"answer": message.content or "No answer returned.", "evidence": last_evidence}

        messages.append(message.model_dump(exclude_none=True))
        evidence = []
        for call in tool_calls:
            arguments = json.loads(call.function.arguments or "{}")
            impl = tool_impls.get(call.function.name)
            try:
                if impl is None:
                    raise ValueError(f"Unknown tool: {call.function.name}")
                output = impl(arguments)
            except (ValueError, KeyError, sqlite3.Error) as error:
                output = {"error": str(error), "rows": [], "row_count": 0}
            evidence.append({"tool": call.function.name, "arguments": arguments, "result": output})
            messages.append({"role": "tool", "tool_call_id": call.id, "content": json.dumps(output, default=str)})
        last_evidence = evidence

    last_errors = []
    if last_evidence:
        for item in last_evidence:
            error = item.get("result", {}).get("error") if isinstance(item.get("result"), dict) else None
            if error:
                last_errors.append(f"{item['tool']}({item['arguments']}): {error}")
    detail = " Last errors: " + "; ".join(last_errors) if last_errors else ""
    return {
        "answer": "I could not complete a verified answer within the tool-call budget." + detail,
        "evidence": last_evidence,
    }


class Handler(BaseHTTPRequestHandler):
    def _send(self, status, payload):
        body = json.dumps(payload).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
        self.end_headers()

    def do_GET(self):
        if self.path == "/api/state":
            return self._send(200, {"workbooks": WORKBOOKS, "activity": ACTIVITY})
        relative_path = self.path.split("?", 1)[0].lstrip("/") or "index.html"
        requested_file = (ROOT / relative_path).resolve()
        if ROOT not in requested_file.parents and requested_file != ROOT:
            return self._send(403, {"error": "Forbidden"})
        if not requested_file.is_file():
            return self._send(404, {"error": "Not found"})
        content_types = {
            ".html": "text/html; charset=utf-8",
            ".js": "text/javascript; charset=utf-8",
            ".css": "text/css; charset=utf-8",
            ".json": "application/json; charset=utf-8",
        }
        body = requested_file.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_types.get(requested_file.suffix, "application/octet-stream"))
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        length = int(self.headers.get("Content-Length", 0))
        try:
            payload = json.loads(self.rfile.read(length))
            if self.path == "/api/workbooks":
                file_bytes = base64.b64decode(payload["content"])
                imported = load_workbook_bytes(file_bytes)
                workbook = {
                    "name": payload.get("name", "workbook.xlsx"),
                    "sheets": f"{len(imported)} sheets",
                    "rows": f"{sum(sheet['rows'] for sheet in imported)} rows",
                    "status": "Synced just now",
                    "tables": [sheet["table"] for sheet in imported],
                }
                WORKBOOKS[:] = [item for item in WORKBOOKS if item["name"] != workbook["name"]]
                WORKBOOKS.append(workbook)
                ACTIVITY.insert(0, {"kind": "Workbook imported", "detail": workbook["name"]})
                return self._send(200, {"name": payload.get("name", "workbook.xlsx"), "sheets": imported})
            if self.path == "/api/ask":
                result = call_groq(payload["question"], payload.get("session_summary", ""))
                ACTIVITY.insert(0, {"kind": "Answer generated", "detail": payload["question"]})
                return self._send(200, result)
            self._send(404, {"error": "Not found"})
        except Exception as error:
            self._send(500, {"error": str(error)})

    def do_DELETE(self):
        length = int(self.headers.get("Content-Length", 0))
        try:
            payload = json.loads(self.rfile.read(length) or b"{}")
            name = str(payload.get("name", "")).strip()
            if not name:
                return self._send(400, {"error": "Workbook name is required"})

            workbook = next((item for item in WORKBOOKS if item["name"] == name), None)
            if workbook is None:
                return self._send(404, {"error": "Workbook not found"})

            for table in workbook.get("tables", []):
                DATABASE.execute(f'DROP TABLE IF EXISTS "{safe_table_name(table)}"')
            DATABASE.commit()
            WORKBOOKS[:] = [item for item in WORKBOOKS if item["name"] != name]
            ACTIVITY.insert(0, {"kind": "Workbook removed", "detail": name})
            return self._send(200, {"name": name})
        except Exception as error:
            self._send(500, {"error": str(error)})

    def log_message(self, format, *args):
        print(format % args)


if __name__ == "__main__":
    print("Ledgerly API listening on http://localhost:8000")
    ThreadingHTTPServer(("localhost", 8000), Handler).serve_forever()